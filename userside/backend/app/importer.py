"""Read contacts from an uploaded .xlsx or .csv file."""
import csv
import io

from openpyxl import load_workbook

PHONE_HEADERS = ("phone", "mobile", "number", "contact", "phone number", "mobile number", "ph", "mob")
NAME_HEADERS = ("name", "customer", "customer name", "full name")
MAX_ROWS = 20000


def _rows(filename: str, data: bytes):
    if filename.lower().endswith(".csv"):
        text = data.decode("utf-8-sig", errors="replace")
        yield from csv.reader(io.StringIO(text))
        return
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        for row in wb.worksheets[0].iter_rows(values_only=True):
            yield list(row)
    finally:
        wb.close()


def _find(header, names):
    for i, cell in enumerate(header):
        if str(cell or "").strip().lower() in names:
            return i
    return None


def read_contacts(filename: str, data: bytes) -> list[tuple[object, str]]:
    """Returns [(raw_phone, name)]. A header row is optional: without one the first
    column is the phone and the second (if any) the name."""
    rows = []
    for row in _rows(filename, data):
        if any(c not in (None, "") for c in row):
            rows.append(row)
        if len(rows) > MAX_ROWS:
            raise ValueError(f"File has more than {MAX_ROWS} rows")
    if not rows:
        return []

    phone_col = _find(rows[0], PHONE_HEADERS)
    name_col = _find(rows[0], NAME_HEADERS)
    if phone_col is not None:
        rows = rows[1:]
    else:
        phone_col, name_col = 0, (1 if len(rows[0]) > 1 else None)

    out = []
    for row in rows:
        phone = row[phone_col] if phone_col < len(row) else None
        name = row[name_col] if name_col is not None and name_col < len(row) else ""
        out.append((phone, str(name or "").strip()[:120]))
    return out
